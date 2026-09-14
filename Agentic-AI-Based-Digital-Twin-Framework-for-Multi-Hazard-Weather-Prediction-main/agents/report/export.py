"""
Export Report, matching the diagram's "Export Report" step. All three
formats are genuinely implemented (not stubs):

- PDF via reportlab (with real matplotlib charts embedded)
- Excel via openpyxl (with native Excel bar/pie charts, not just images)
- JSON: the sections list as-is

"dashboard" format returns the same structured JSON a frontend would
consume to render its own dashboard -- there's no separate dashboard
renderer here, that's frontend work.
"""

from __future__ import annotations

import io
import json
import os
from typing import Any, Dict, List

from .logging_config import get_logger

logger = get_logger(__name__)


def _render_chart_png(chart: Dict[str, Any]) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(6, 3.5))
    data = chart["content"]["data"]
    if chart["content"]["type"] == "bar":
        ax.bar(list(data.keys()), list(data.values()), color="#3498db")
        ax.set_ylabel(chart["content"].get("ylabel", ""))
    elif chart["content"]["type"] == "pie":
        ax.pie(list(data.values()), labels=list(data.keys()), autopct="%1.0f%%")
    ax.set_title(chart["title"])
    fig.tight_layout()

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=120)
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def export_pdf(sections: List[Dict[str, Any]], metadata: Dict[str, Any], output_path: str) -> str:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(output_path, pagesize=A4)
    story = [
        Paragraph(f"{metadata['report_type'].replace('_', ' ').title()} Report", styles["Title"]),
        Paragraph(f"Districts: {', '.join(metadata['districts'])} | Generated: {metadata['generated_at']} "
                  f"| Version: {metadata['version']}", styles["Normal"]),
        Spacer(1, 0.2 * inch),
    ]

    for section in sections:
        story.append(Paragraph(section["title"], styles["Heading2"]))
        if section["kind"] == "text":
            story.append(Paragraph(section["content"].replace("\n", "<br/>"), styles["Normal"]))
        elif section["kind"] == "table":
            rows = section["content"]
            if rows:
                headers = list(rows[0].keys())
                table_data = [headers] + [[str(r.get(h, "")) for h in headers] for r in rows]
                table = Table(table_data, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3498db")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]))
                story.append(table)
        elif section["kind"] == "chart":
            png_bytes = _render_chart_png(section)
            story.append(Image(io.BytesIO(png_bytes), width=5 * inch, height=2.9 * inch))
        story.append(Spacer(1, 0.15 * inch))

    doc.build(story)
    logger.info("Exported PDF report to %s", output_path)
    return output_path


def export_excel(sections: List[Dict[str, Any]], metadata: Dict[str, Any], output_path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.styles import Font

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    summary_ws["A1"] = f"{metadata['report_type'].replace('_', ' ').title()} Report"
    summary_ws["A1"].font = Font(bold=True, size=14)
    summary_ws["A2"] = f"Districts: {', '.join(metadata['districts'])}"
    summary_ws["A3"] = f"Generated: {metadata['generated_at']} | Version: {metadata['version']}"

    row = 5
    for section in sections:
        if section["kind"] != "text":
            continue
        summary_ws.cell(row=row, column=1, value=section["title"]).font = Font(bold=True)
        row += 1
        for line in section["content"].split("\n"):
            summary_ws.cell(row=row, column=1, value=line)
            row += 1
        row += 1

    for section in sections:
        if section["kind"] == "table" and section["content"]:
            ws = wb.create_sheet(title=section["title"][:31])
            headers = list(section["content"][0].keys())
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for record in section["content"]:
                ws.append([record.get(h) for h in headers])

        elif section["kind"] == "chart":
            ws = wb.create_sheet(title=section["title"][:31])
            data = section["content"]["data"]
            ws.append(["Label", "Value"])
            for k, v in data.items():
                ws.append([k, v])

            chart = BarChart() if section["content"]["type"] == "bar" else PieChart()
            chart.title = section["title"]
            values_ref = Reference(ws, min_col=2, min_row=1, max_row=len(data) + 1)
            cats_ref = Reference(ws, min_col=1, min_row=2, max_row=len(data) + 1)
            chart.add_data(values_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            ws.add_chart(chart, "D2")

    wb.save(output_path)
    logger.info("Exported Excel report to %s", output_path)
    return output_path


def export_json(sections: List[Dict[str, Any]], metadata: Dict[str, Any], output_path: str) -> str:
    with open(output_path, "w") as f:
        json.dump({"metadata": metadata, "sections": sections}, f, indent=2, default=str)
    logger.info("Exported JSON report to %s", output_path)
    return output_path


def export_report(sections: List[Dict[str, Any]], metadata: Dict[str, Any], fmt: str,
                   output_dir: str) -> Dict[str, Any]:
    os.makedirs(output_dir, exist_ok=True)
    base_name = f"{metadata['report_id']}"

    if fmt == "pdf":
        path = export_pdf(sections, metadata, os.path.join(output_dir, f"{base_name}.pdf"))
    elif fmt == "excel":
        path = export_excel(sections, metadata, os.path.join(output_dir, f"{base_name}.xlsx"))
    elif fmt in ("json", "dashboard"):
        path = export_json(sections, metadata, os.path.join(output_dir, f"{base_name}.json"))
    else:
        raise ValueError(f"Unsupported format: {fmt!r}")

    return {"file_path": path, "format": fmt}
