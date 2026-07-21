import ast
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List


@dataclass
class FileContent:
    text: str
    metadata: Dict[str, object]


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_json(path: Path) -> str:
    try:
        obj = json.loads(_read_text(path))
        return json.dumps(obj, indent=2, ensure_ascii=False)
    except Exception:
        return _read_text(path)


def _read_yaml(path: Path) -> str:
    try:
        import yaml

        obj = yaml.safe_load(_read_text(path))
        return yaml.safe_dump(obj, allow_unicode=True, sort_keys=False)
    except Exception:
        return _read_text(path)


def _read_ipynb(path: Path) -> str:
    try:
        notebook = json.loads(_read_text(path))
        cells = []
        for cell in notebook.get("cells", []):
            cell_type = cell.get("cell_type")
            source = "".join(cell.get("source", []))
            if cell_type == "markdown":
                cells.append(source)
            elif cell_type == "code":
                cells.append("# Code cell:\n" + source)
        return "\n\n".join(cells)
    except Exception:
        return _read_text(path)


def _read_python(path: Path) -> str:
    text = _read_text(path)
    comments = re.findall(r"^\s*#(.*)$", text, flags=re.MULTILINE)
    docstrings: List[str] = []
    try:
        tree = ast.parse(text)
        for node in ast.walk(tree):
            if isinstance(node, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                docstring = ast.get_docstring(node)
                if docstring:
                    docstrings.append(docstring)
    except Exception:
        pass
    content = [f"# File: {path.name}"]
    if docstrings:
        content.append("\n\n".join(docstrings))
    if comments:
        content.append("\n\n".join(comments[:500]))
    return "\n\n".join(content) or text


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ImportError("DOCX support requires python-docx. Install it and retry.") from exc

    document = docx.Document(path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _read_pptx(path: Path) -> str:
    try:
        import pptx
    except ImportError as exc:
        raise ImportError("PPTX support requires python-pptx. Install it and retry.") from exc

    presentation = pptx.Presentation(path)
    slides = []
    for slide in presentation.slides:
        shapes = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        if shapes:
            slides.append("\n".join(shapes))
    return "\n\n".join(slides)


def _read_odt(path: Path) -> str:
    try:
        from odf import text, teletype
        from odf.opendocument import load
    except ImportError as exc:
        raise ImportError("ODT support requires odfpy. Install it and retry.") from exc

    document = load(str(path))
    paragraphs = []
    for element in document.getElementsByType(text.P):
        paragraphs.append(teletype.extractText(element))
    return "\n\n".join([p for p in paragraphs if p.strip()])


def _read_rtf(path: Path) -> str:
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError as exc:
        raise ImportError("RTF support requires striprtf. Install it and retry.") from exc

    return rtf_to_text(_read_text(path))


def _read_pdf(path: Path) -> str:
    try:
        from RAG_project.loaders.pdf_loader import PDFLoader
    except Exception as exc:
        raise ImportError("PDF support requires the existing RAG_project PDF loader and its dependencies.") from exc

    loader = PDFLoader(path.parent)
    pages = loader.load_pdf(path)
    texts = [page.page_content for page in pages]
    return "\n\n".join(texts)


def read_file(path: Path) -> FileContent:
    suffix = path.suffix.lower()
    metadata = {
        "file_name": path.name,
        "file_size": path.stat().st_size,
        "file_extension": suffix,
    }
    if suffix in {".txt", ".md", ".markdown", ".rst"}:
        return FileContent(text=_read_text(path), metadata=metadata)
    if suffix == ".json":
        return FileContent(text=_read_json(path), metadata=metadata)
    if suffix in {".yaml", ".yml"}:
        return FileContent(text=_read_yaml(path), metadata=metadata)
    if suffix == ".ipynb":
        return FileContent(text=_read_ipynb(path), metadata=metadata)
    if suffix == ".py":
        return FileContent(text=_read_python(path), metadata=metadata)
    if suffix == ".docx":
        return FileContent(text=_read_docx(path), metadata=metadata)
    if suffix == ".pptx":
        return FileContent(text=_read_pptx(path), metadata=metadata)
    if suffix == ".odt":
        return FileContent(text=_read_odt(path), metadata=metadata)
    if suffix == ".rtf":
        return FileContent(text=_read_rtf(path), metadata=metadata)
    if suffix == ".pdf":
        return FileContent(text=_read_pdf(path), metadata=metadata)
    if path.name.lower() == "dockerfile" or path.name.lower() == "requirements.txt":
        return FileContent(text=_read_text(path), metadata=metadata)
    return FileContent(text=_read_text(path), metadata=metadata)
